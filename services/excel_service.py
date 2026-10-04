"""Excel integration (outside the browser)."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


class ExcelService:
    def __init__(self, output_dir: Path, process_name: str) -> None:
        self._output_dir = output_dir
        self._process_name = process_name

    def write_rows(self, headers: Sequence[str], rows: Sequence[Sequence[Any]], sheet_name: str = "Data") -> Path:
        """Write a new workbook named <process>_<YYYYMMDD_HHMMSS>.xlsx and return its path."""
        self._output_dir.mkdir(parents=True, exist_ok=True)
        output_file = self._output_dir / f"{self._process_name}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(list(headers))
        for row in rows:
            sheet.append(list(row))
        workbook.save(output_file)
        return output_file

    def read_rows(self, input_file: Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
        """Read a sheet whose first row is the header. Returns one dict per row."""
        workbook = load_workbook(input_file, read_only=True, data_only=True)
        try:
            sheet = workbook[sheet_name] if sheet_name else workbook.active
            values = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not values:
            return []
        headers = [str(header) for header in values[0]]
        return [dict(zip(headers, row)) for row in values[1:]]
